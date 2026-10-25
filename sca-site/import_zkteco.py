#!/usr/bin/env python3
"""
import_zkteco.py - importe les produits ZKTeco (products.xlsx) dans le site.

Ce que fait le script :
  1) database/data/products.json  : ajoute les 127 produits (clé = slug)
  2) database/data/solutions.json : remplit la catégorie
        Sécurité électronique > « Contrôle d'accès électronique »
        (slug : terminaux-de-reconnaissance-faciale)
  3) Aucun prix n'est importé (le site n'affiche que « Demande devis »).

Sécurité :
  - sauvegarde automatique (.bak-AAAAMMJJ-HHMM) avant toute écriture
  - relançable sans danger : les produits ZKTeco sont remplacés, rien d'autre n'est touché
  - --dry-run : affiche le résultat sans écrire

Utilisation (depuis n'importe où) :
    python import_zkteco.py --xlsx products.xlsx --project C:/chemin/vers/laravel --dry-run
    python import_zkteco.py --xlsx products.xlsx --project C:/chemin/vers/laravel
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

SOLUTION_SLUG = "securite-electroniques"
CATEGORY_SLUG = "terminaux-de-reconnaissance-faciale"
BRAND = "ZKTeco"

# catégorie du PDF -> (ordre, nom du groupe affiché)
GROUPS = {
    "Terminal de Contrôle d'Accès": (1, "Terminaux"),
    "Lecteur de Carte & Tag ID/IC": (2, "Lecteurs"),
    "Panneau de Contrôle d'Accès": (3, "Panneaux de contrôle"),
}

# traduction des libellés de spécifications les plus fréquents
LABELS = {
    "Communication": "Communication",
    "Power Supply": "Alimentation",
    "Operating Temperature": "Température de fonctionnement",
    "Operating Humidity": "Humidité de fonctionnement",
    "Storage Temperature": "Température de stockage",
    "Dimension": "Dimensions",
    "Dimensions": "Dimensions",
    "Dimensions (L × W × H)": "Dimensions (L × l × H)",
    "Access Control Interface": "Interface de contrôle d'accès",
    "Hardware": "Matériel",
    "Card Capacity": "Capacité de cartes",
    "Card Holders Capacity": "Capacité de porteurs de cartes",
    "Transaction Capacity": "Capacité de transactions",
    "Net Weight": "Poids net",
    "Gross Weight": "Poids brut",
    "Weight": "Poids",
    "Package Weight": "Poids du colis",
    "Certifications": "Certifications",
    "Ingress Protection Rating": "Indice de protection (IP)",
    "Standard Functions": "Fonctions standard",
    "Optional Functions": "Fonctions optionnelles",
    "Display": "Écran",
    "Installation": "Installation",
    "Mounting": "Montage",
    "Model": "Modèle",
    "Fingerprint Capacity": "Capacité d'empreintes",
    "Fingerprint Template Capacity": "Capacité de modèles d'empreintes",
    "Face Capacity": "Capacité de visages",
    "Palm Capacity": "Capacité de paumes",
    "User Capacity": "Capacité d'utilisateurs",
    "Log Events Capacity": "Capacité d'événements",
    "Authentication Method": "Méthode d'authentification",
    "Card Type": "Type de carte",
    "Supported Software": "Logiciels compatibles",
    "Operation System": "Système d'exploitation",
    "Operating System": "Système d'exploitation",
    "Biometric Algorithm": "Algorithme biométrique",
    "Biometric Verification Speed": "Vitesse de vérification biométrique",
    "Verification Speed": "Vitesse de vérification",
    "Specification": "Spécification",
    "Housing Material": "Matériau du boîtier",
    "Enclosure": "Boîtier",
    "Reading Distance": "Distance de lecture",
    "Recognition Distance": "Distance de reconnaissance",
    "Number of Readers Supported": "Nombre de lecteurs pris en charge",
    "Types of Readers Supported": "Types de lecteurs pris en charge",
    "Number of Inputs": "Nombre d'entrées",
    "Number of Outputs": "Nombre de sorties",
    "Number of Doors Controlled": "Nombre de portes contrôlées",
    "Wiegand Signal": "Signal Wiegand",
    "Operating Environment": "Environnement de fonctionnement",
    "Operating Voltage": "Tension de fonctionnement",
    "Current Draw": "Consommation",
    "Tamper Switch": "Contact anti-arrachement",
    "Visual Indicator": "Indicateur visuel",
    "Audio Indicator": "Indicateur sonore",
    "Physical Keypad": "Clavier physique",
    "Compatibility": "Compatibilité",
    "Flash Memory": "Mémoire flash",
    "Camera": "Caméra",
    "Algorithm Version": "Version de l'algorithme",
}


def norm_label(label: str) -> str:
    s = str(label or "").strip()
    s = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", s)          # PowerSupply -> Power Supply
    s = re.sub(r"\bNumberof\b", "Number of", s)
    s = re.sub(r"\bTypesof\b", "Types of", s)
    s = re.sub(r"\s+", " ", s)
    return LABELS.get(s, s)


def clean_description(text: str) -> tuple[str, bool]:
    """Retire la fin tronquée '…'. Retourne (texte, était_tronqué)."""
    t = re.sub(r"\s+", " ", (text or "")).strip()
    if not t.endswith("…"):
        return t, False
    t = t[:-1].rstrip()
    cut = max(t.rfind(". "), t.rfind(".\u00a0"))
    if cut >= 50:
        return t[: cut + 1], True                        # dernière phrase complète
    cut = t.rfind(" ")
    return (t[:cut] if cut > 0 else t).rstrip(" ,;:-"), True


def short_lead(text: str, limit: int = 140) -> str:
    if len(text) <= limit:
        return text
    cut = text.rfind(" ", 0, limit)
    return text[: cut if cut > 0 else limit].rstrip(" ,;:-") + "…"


def natural_key(s: str):
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r"(\d+)", s or "")]


def load_json(path: Path, default):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def backup(path: Path, stamp: str):
    if path.exists():
        shutil.copy2(path, path.with_name(path.name + f".bak-{stamp}"))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default="products.xlsx")
    ap.add_argument("--project", required=True, help="racine du projet Laravel")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-check-images", action="store_true",
                    help="n'exige pas que les images existent déjà dans public/images/products")
    a = ap.parse_args()

    root = Path(a.project)
    data_dir = root / "database" / "data"
    public = root / "public"
    products_file = data_dir / "products.json"
    solutions_file = data_dir / "solutions.json"
    if not solutions_file.exists():
        sys.exit(f"Introuvable : {solutions_file}")

    # ---------- lecture Excel ----------
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    head = [c.value for c in next(wb["products"].iter_rows(min_row=1, max_row=1))]
    rows = [dict(zip(head, [c.value for c in r])) for r in wb["products"].iter_rows(min_row=2)]
    rows = [r for r in rows if r.get("slug")]

    specs_by_slug: dict[str, list[tuple[str, str]]] = {}
    for slug, _ordre, label, value in list(wb["specs"].iter_rows(min_row=2, values_only=True)):
        if slug and label and value not in (None, ""):
            specs_by_slug.setdefault(slug, []).append((label, str(value).strip()))

    # ---------- état existant ----------
    existing = load_json(products_file, {})
    if isinstance(existing, list):
        sys.exit("products.json est une liste : format inattendu (le site attend une clé = slug).")

    new_products: dict[str, dict] = {}
    truncated = missing_img = no_img = renamed = 0

    for r in rows:
        src_slug = str(r["slug"]).strip().lower()
        slug = src_slug
        if slug in existing and existing[slug].get("brand") != BRAND:
            slug = "zkteco-" + src_slug                  # évite d'écraser un produit existant
            renamed += 1

        order, group = GROUPS.get(r.get("categorie"), (9, "Autres"))

        # images : output/images/products/<slug>/<slug>-N.png  ->  images/products/<slug>/<slug>-N.webp
        gallery = []
        for p in str(r.get("images") or "").split(" | "):
            name = Path(p.strip()).stem
            if not name:
                continue
            rel = f"images/products/{src_slug}/{name}.webp"
            if a.no_check_images or (public / rel).exists():
                gallery.append(rel)
            else:
                missing_img += 1
        if not gallery:
            no_img += 1

        desc, was_cut = clean_description(r.get("description") or "")
        truncated += was_cut

        seen, specs = set(), []
        for label, value in specs_by_slug.get(src_slug, []):
            lab = norm_label(label)
            if lab.lower() in seen:
                continue
            seen.add(lab.lower())
            specs.append({"label": lab, "value": value})

        new_products[slug] = {
            "slug": slug,
            "title": (r.get("nom") or src_slug).strip(),
            "subtitle": desc,
            "lead": short_lead(desc),
            "brand": BRAND,
            "group": group,
            "group_order": order,
            "image": gallery[0] if gallery else None,
            "gallery": gallery,
            "features": [],
            "paragraphs": [],
            "specs": specs,
            "pdfs": [],
            "solution": SOLUTION_SLUG,
            "category": CATEGORY_SLUG,
        }

    # ordre d'affichage : groupe puis nom
    ordered = sorted(new_products.values(),
                     key=lambda p: (p["group_order"], natural_key(p["title"])))
    for p in ordered:
        p.pop("group_order")
    ordered_slugs = [p["slug"] for p in ordered]

    # ---------- fusion products.json ----------
    merged = {k: v for k, v in existing.items() if v.get("brand") != BRAND}
    for p in ordered:
        merged[p["slug"]] = p

    # ---------- solutions.json ----------
    solutions = load_json(solutions_file, [])
    sol = next((s for s in solutions if s.get("slug") == SOLUTION_SLUG), None)
    cat = next((c for c in (sol or {}).get("categories", []) if c.get("slug") == CATEGORY_SLUG), None)
    if not cat:
        sys.exit(f"Catégorie {SOLUTION_SLUG}/{CATEGORY_SLUG} introuvable dans solutions.json")
    keep = [s for s in cat.get("products", []) if merged.get(s, {}).get("brand") != BRAND and s in merged]
    cat["products"] = keep + ordered_slugs

    # ---------- rapport ----------
    by_group = {}
    for p in ordered:
        by_group[p["group"]] = by_group.get(p["group"], 0) + 1
    print(f"Produits lus            : {len(rows)}")
    print(f"Par groupe              : {by_group}")
    print(f"Sans image              : {no_img}")
    print(f"Images manquantes       : {missing_img}  (lancez d'abord prepare_images.py)")
    print(f"Descriptions tronquées  : {truncated}  (nettoyées ; version complète à récupérer)")
    print(f"Slugs renommés (conflit): {renamed}")
    print(f"products.json           : {len(existing)} -> {len(merged)} produits")
    print(f"Catégorie {CATEGORY_SLUG} : {len(cat['products'])} produits")

    if a.dry_run:
        print("\n--dry-run : rien n'a été écrit.")
        return

    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    data_dir.mkdir(parents=True, exist_ok=True)
    backup(products_file, stamp)
    backup(solutions_file, stamp)
    products_file.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")
    solutions_file.write_text(json.dumps(solutions, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nÉcrit. Sauvegardes : *.bak-{stamp}")


if __name__ == "__main__":
    main()